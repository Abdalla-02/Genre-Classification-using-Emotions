"""Export docs/current_state.md to docs/current_state.docx and docs/current_state.pdf.

The supervisor briefing is written in Markdown; this turns it into a Word document with
python-docx (headings, paragraphs, bullet and numbered lists, tables, bold / italic /
code, embedded images) and then into PDF with LibreOffice. It covers the Markdown subset
the briefing uses, not Markdown in general.

Run:  python experiments/export_current_state.py
      (needs LibreOffice for the PDF; without it only the .docx is written)
"""

from __future__ import annotations

import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DOCS = Path(__file__).resolve().parents[1] / "docs"
SRC = DOCS / "current_state.md"
SOFFICE_CANDIDATES = ["soffice", r"C:\Program Files\LibreOffice\program\soffice.exe"]

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)")


def add_runs(par, text, bold=False):
    """Inline Markdown: **bold**, *italic*, `code`."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = par.add_run(part[1:-1]); r.italic = True
        else:
            r = par.add_run(part)
        if bold:
            r.bold = True


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    header, align, body = cells[0], cells[1], cells[2:]
    right = [a.endswith(":") and not a.startswith(":") for a in align]
    t = doc.add_table(rows=1 + len(body), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate([header] + body):
        for j, txt in enumerate(row[:len(header)]):
            cell = t.cell(i, j)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_runs(p, txt, bold=(i == 0))
            for r in p.runs:
                r.font.size = Pt(8.5)
            if right[j] and i > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if i == 0:
                shade(cell, "E8E8E8")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def convert(md: str, out: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"; normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)

    lines = md.split("\n")
    i, para = 0, []

    def flush():
        if para:
            add_runs(doc.add_paragraph(), " ".join(s.strip() for s in para))
            para.clear()

    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            flush(); i += 1; continue
        if s == "---":
            flush(); i += 1; continue
        m = re.match(r"(#{1,4}) (.*)", s)
        if m:
            flush()
            doc.add_heading(m.group(2).replace("`", ""), level=min(len(m.group(1)) - 1, 3) or 0)
            i += 1; continue
        if s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i]); i += 1
            add_table(doc, rows); continue
        m = re.match(r"!\[(.*?)\]\((.*?)\)", s)
        if m:
            flush()
            img = (SRC.parent / m.group(2)).resolve()
            if img.exists():
                doc.add_picture(str(img), width=Cm(15.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph(); add_runs(cap, f"*{m.group(1)}*")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1; continue
        if s.startswith("```"):
            flush(); i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                p = doc.add_paragraph(); r = p.add_run(lines[i])
                r.font.name = "Consolas"; r.font.size = Pt(8)
                p.paragraph_format.space_after = Pt(0)
                i += 1
            i += 1; continue
        m = re.match(r"( *)([-*]|\d+\.) (.*)", line)
        if m:
            flush()
            indent = len(m.group(1)) // 2
            numbered = m.group(2)[0].isdigit()
            text = [m.group(3)]
            i += 1
            # continuation lines of the same item
            while i < len(lines) and lines[i].startswith(" " * (len(m.group(1)) + 2)) \
                    and not re.match(r" *([-*]|\d+\.) ", lines[i]) and lines[i].strip():
                text.append(lines[i].strip()); i += 1
            if numbered:
                # literal numbers: Word's automatic numbering would run on across lists
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.9 + 0.6 * indent)
                p.paragraph_format.first_line_indent = Cm(-0.6)
                add_runs(p, m.group(2) + " " + " ".join(text))
            else:
                p = doc.add_paragraph(style="List Bullet 2" if indent else "List Bullet")
                if indent:
                    p.paragraph_format.left_indent = Cm(1.6)
                add_runs(p, " ".join(text))
            p.paragraph_format.space_after = Pt(2)
            continue
        para.append(line)
        i += 1
    flush()
    doc.save(out)


def main() -> None:
    md = SRC.read_text(encoding="utf-8")
    docx_path = SRC.with_suffix(".docx")
    convert(md, docx_path)
    print(f"wrote {docx_path}")
    soffice = next((c for c in SOFFICE_CANDIDATES if shutil.which(c) or Path(c).exists()), None)
    if not soffice:
        sys.exit("LibreOffice not found: PDF not written")
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(DOCS),
                    str(docx_path)], check=True, capture_output=True)
    print(f"wrote {SRC.with_suffix('.pdf')}")


if __name__ == "__main__":
    main()
